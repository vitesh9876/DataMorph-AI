import logging
from docx import Document
import pandas as pd
from app.services.extractors.base import ExtractionResult
from app.services.extractors.unstructured_structurer import UnstructuredDataStructurer

logger = logging.getLogger(__name__)

class DOCXExtractor:
    """Extracts tables, paragraphs, headers, and structures free-form text and data from Word DOCX files."""

    @staticmethod
    def extract_docx(file_path: str) -> ExtractionResult:
        doc = Document(file_path)
        
        paragraphs_text = []
        headings_count = 0
        
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs_text.append(text)
                if p.style and p.style.name.startswith("Heading"):
                    headings_count += 1

        extracted_tables = []
        for table in doc.tables:
            rows_data = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_cells)

            if rows_data and len(rows_data) > 1:
                header = [f"Col_{i+1}" if not col else col for i, col in enumerate(rows_data[0])]
                df_table = pd.DataFrame(rows_data[1:], columns=header)
                extracted_tables.append(df_table)

        full_raw_text = "\n\n".join(paragraphs_text)

        if extracted_tables:
            primary_df = extracted_tables[0]
            data_type = "semi-structured"
        else:
            primary_df, meta, sections = UnstructuredDataStructurer.structure_unstructured_text(full_raw_text)
            data_type = "unstructured_document"
            headings_count = max(headings_count, meta.get("sections_count", 1))

        return ExtractionResult(
            dataframe=primary_df,
            raw_text=full_raw_text[:8000],
            tables=extracted_tables if extracted_tables else [primary_df],
            doc_stats={
                "pages": max(len(paragraphs_text) // 10, 1),
                "tables": max(len(extracted_tables), 1),
                "charts": 0,
                "sections": max(headings_count, 1),
                "rows": int(len(primary_df)),
                "columns": int(len(primary_df.columns))
            },
            data_type=data_type,
            metadata={"paragraphs_count": len(paragraphs_text)}
        )
