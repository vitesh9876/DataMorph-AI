import io
import os
import re
from typing import List, Dict, Any
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

class PPTXExporter:
    """Exports structured reports to high-impact widescreen 16:9 PowerPoint (.pptx) presentations."""

    @classmethod
    def _create_chart_image(cls, chart_type: str, chart_data: List[Dict[str, Any]]) -> io.BytesIO:
        """Generates a high-resolution chart figure for embedding in a PowerPoint slide."""
        fig, ax = plt.subplots(figsize=(6.5, 4.2), dpi=180)
        fig.patch.set_facecolor("#FFFFFF")
        ax.set_facecolor("#F8FAFC")

        items = chart_data[:8]
        labels = [str(d.get("name") or d.get("label") or d.get("category") or f"Item {i+1}") for i, d in enumerate(items)]
        values = []
        for d in items:
            val = d.get("value") or d.get("count") or 0
            try:
                values.append(float(val))
            except (ValueError, TypeError):
                values.append(0.0)

        palette = ["#6366F1", "#3B82F6", "#10B981", "#F59E0B", "#EF4444", "#8B5CF6", "#EC4899", "#14B8A6"]

        c_type = chart_type.lower()
        if "pie" in c_type or "donut" in c_type:
            wedges, texts, autotexts = ax.pie(
                values,
                labels=labels,
                autopct="%1.1f%%",
                startangle=140,
                colors=palette[:len(values)],
                wedgeprops=dict(width=0.45 if "donut" in c_type else 1.0, edgecolor="#FFFFFF", linewidth=2)
            )
            for at in autotexts:
                at.set_fontsize(8)
                at.set_color("#1E293B")
                at.set_weight("bold")
            for t in texts:
                t.set_fontsize(8)
                t.set_color("#475569")
        elif "line" in c_type:
            ax.plot(labels, values, color="#6366F1", marker="o", linewidth=2.5, markersize=5)
            ax.fill_between(range(len(values)), values, color="#6366F1", alpha=0.15)
            ax.grid(True, linestyle="--", alpha=0.4, color="#CBD5E1")
            ax.tick_params(axis="x", rotation=25, labelsize=8)
            ax.tick_params(axis="y", labelsize=8)
        else:
            bars = ax.bar(labels, values, color=palette[:len(values)], width=0.55, edgecolor="#E2E8F0")
            ax.grid(axis="y", linestyle="--", alpha=0.4, color="#CBD5E1")
            ax.tick_params(axis="x", rotation=25, labelsize=8)
            ax.tick_params(axis="y", labelsize=8)
            for bar in bars:
                height = bar.get_height()
                ax.annotate(
                    f"{height:,.0f}" if height >= 10 else f"{height:,.2f}",
                    xy=(bar.get_x() + bar.get_width() / 2, height),
                    xytext=(0, 3),
                    textcoords="offset points",
                    ha="center",
                    va="bottom",
                    fontsize=7,
                    fontweight="bold",
                    color="#1E293B"
                )

        for spine in ["top", "right"]:
            ax.spines[spine].set_visible(False)
        for spine in ["left", "bottom"]:
            ax.spines[spine].set_color("#CBD5E1")

        plt.tight_layout()
        buf = io.BytesIO()
        plt.savefig(buf, format="png", bbox_inches="tight", dpi=180, facecolor=fig.get_facecolor())
        plt.close(fig)
        buf.seek(0)
        return buf

    @classmethod
    def export_to_pptx(
        cls,
        output_path: str,
        report_title: str,
        template_type: str = "professional",
        sections: List[Dict[str, Any]] = None
    ) -> str:
        if sections is None:
            sections = []

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_slide_layout = prs.slide_layouts[6]

        # 1. Title Slide
        title_slide = prs.slides.add_slide(blank_slide_layout)
        
        txBox = title_slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(3.5))
        tf = txBox.text_frame
        tf.word_wrap = True

        p_badge = tf.paragraphs[0]
        p_badge.text = "DATAMORPH AI • EXECUTIVE INTELLIGENCE BRIEF"
        p_badge.font.size = Pt(13)
        p_badge.font.bold = True
        p_badge.font.color.rgb = RGBColor(99, 102, 241)

        p_title = tf.add_paragraph()
        p_title.text = report_title
        p_title.font.size = Pt(34)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(15, 23, 42)
        p_title.space_before = Pt(12)

        p_sub = tf.add_paragraph()
        p_sub.text = "Automated Data Structuring, Machine Learning Analytics & Verified Visual Studio"
        p_sub.font.size = Pt(15)
        p_sub.font.color.rgb = RGBColor(100, 116, 139)
        p_sub.space_before = Pt(8)

        # 2. Section Slides
        for s in sections:
            if not s.get("is_visible", True) or s.get("section_type") == "title_page":
                continue

            sec_title = s.get("title", "")
            raw_content = s.get("content", "")
            sec_type = s.get("section_type", "")

            slide = prs.slides.add_slide(blank_slide_layout)

            # Slide Header
            header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
            htf = header_box.text_frame
            hp = htf.paragraphs[0]
            hp.text = sec_title
            hp.font.size = Pt(22)
            hp.font.bold = True
            hp.font.color.rgb = RGBColor(30, 58, 138)

            # Case A: Visual Chart Section
            if sec_type == "visualizations" and s.get("chart_data"):
                chart_type = s.get("chart_type", "bar")
                chart_data = s.get("chart_data", [])

                try:
                    img_buf = cls._create_chart_image(chart_type, chart_data)
                    slide.shapes.add_picture(img_buf, Inches(0.8), Inches(1.5), width=Inches(6.8))
                except Exception:
                    pass

                # Insight Box on the right side
                desc_box = slide.shapes.add_textbox(Inches(8.0), Inches(1.5), Inches(4.5), Inches(5.2))
                dtf = desc_box.text_frame
                dtf.word_wrap = True

                dp_badge = dtf.paragraphs[0]
                dp_badge.text = f"ARCHETYPE: {chart_type.upper().replace('_', ' ')}"
                dp_badge.font.size = Pt(11)
                dp_badge.font.bold = True
                dp_badge.font.color.rgb = RGBColor(99, 102, 241)

                clean_text = raw_content.replace("**", "").split("Data Breakdown:")[0].strip()
                dp_body = dtf.add_paragraph()
                dp_body.text = clean_text
                dp_body.font.size = Pt(13)
                dp_body.font.color.rgb = RGBColor(51, 65, 85)
                dp_body.space_before = Pt(8)
                continue

            # Case B: Structured Table Section
            if sec_type == "structured_table":
                table_data = s.get("table_data") or []
                table_columns = s.get("table_columns") or []
                if not table_columns and table_data:
                    table_columns = list(table_data[0].keys())

                if table_data and table_columns:
                    disp_cols = table_columns[:6]
                    disp_rows = table_data[:12]

                    rows_count = len(disp_rows) + 1
                    cols_count = len(disp_cols)

                    table_shape = slide.shapes.add_table(
                        rows_count, cols_count, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.4 * rows_count)
                    )
                    table = table_shape.table

                    # Header Row
                    for col_idx, col_name in enumerate(disp_cols):
                        cell = table.cell(0, col_idx)
                        cell.text = str(col_name)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(99, 102, 241)
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = Pt(11)
                            paragraph.font.bold = True
                            paragraph.font.color.rgb = RGBColor(255, 255, 255)

                    # Data Rows
                    for row_idx, row_dict in enumerate(disp_rows):
                        for col_idx, col_name in enumerate(disp_cols):
                            cell = table.cell(row_idx + 1, col_idx)
                            cell.text = str(row_dict.get(col_name, "—"))
                            cell.fill.solid()
                            if row_idx % 2 == 0:
                                cell.fill.fore_color.rgb = RGBColor(248, 250, 252)
                            else:
                                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.size = Pt(10)
                                paragraph.font.color.rgb = RGBColor(30, 41, 59)
                    continue

            # Case C: Standard Text / Executive Summary Slide
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
            ctf = content_box.text_frame
            ctf.word_wrap = True

            paragraphs = raw_content.split("\n")
            first_p = True
            for line in paragraphs:
                trimmed = line.strip().replace("**", "")
                if not trimmed:
                    continue

                if first_p:
                    cp = ctf.paragraphs[0]
                    first_p = False
                else:
                    cp = ctf.add_paragraph()

                cp.text = trimmed
                cp.font.size = Pt(14)
                cp.font.color.rgb = RGBColor(51, 65, 85)
                cp.space_before = Pt(6)

                if trimmed.startswith("- ") or trimmed.startswith("* "):
                    cp.level = 1
                    cp.text = trimmed[2:].strip()

        prs.save(output_path)
        return output_path
