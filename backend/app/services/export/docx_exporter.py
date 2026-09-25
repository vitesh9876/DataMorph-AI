import os
import re
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class DOCXExporter:
    """Exports structured reports to Microsoft Word (.docx) with DataMorph AI header branding."""

    @classmethod
    def export_to_docx(
        cls,
        output_path: str,
        report_title: str,
        template_type: str,
        sections: List[Dict[str, Any]]
    ) -> str:
        doc = Document()
        
        # Margins
        for sec in doc.sections:
            sec.top_margin = Inches(0.8)
            sec.bottom_margin = Inches(0.8)
            sec.left_margin = Inches(0.8)
            sec.right_margin = Inches(0.8)

        # Brand Tag Header
        brand_p = doc.add_paragraph()
        brand_p.paragraph_format.space_before = Pt(0)
        brand_p.paragraph_format.space_after = Pt(2)
        r_brand = brand_p.add_run("DATAMORPH AI • VISUAL INTELLIGENCE WORKSPACE")
        r_brand.font.name = "Arial"
        r_brand.font.size = Pt(9.5)
        r_brand.font.bold = True
        r_brand.font.color.rgb = RGBColor(99, 102, 241)

        # Title
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(2)
        run_title = title_p.add_run(report_title)
        run_title.font.name = "Arial"
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        # Subtitle
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(14)
        run_sub = sub_p.add_run("Automated Data Structuring & Verified Visual Analytics")
        run_sub.font.size = Pt(9)
        run_sub.font.color.rgb = RGBColor(100, 116, 139)

        # Iterate Sections
        for s in sections:
            if not s.get("is_visible", True):
                continue

            if s.get("section_type") == "title_page":
                continue

            sec_title = s.get("title", "")
            raw_content = s.get("content", "")

            # Heading
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            r_h = h.add_run(sec_title)
            r_h.font.name = "Arial"
            r_h.font.size = Pt(13)
            r_h.font.bold = True
            r_h.font.color.rgb = RGBColor(30, 41, 59)

            # Content paragraphs
            lines = raw_content.split("\n")
            for line in lines:
                clean_line = line.strip()
                if not clean_line:
                    continue

                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15

                # Strip markdown bold or format
                clean_text = clean_line.replace("**", "")
                r_body = p.add_run(clean_text)
                r_body.font.name = "Calibri"
                r_body.font.size = Pt(10)
                r_body.font.color.rgb = RGBColor(51, 65, 85)

        doc.save(output_path)
        return output_path
