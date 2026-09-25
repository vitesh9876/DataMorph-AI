import os
from typing import List, Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class XLSXExporter:
    """Exports cleaned structured datasets and audit metrics to styled multi-tab Excel workbooks."""

    @classmethod
    def export_to_xlsx(
        cls,
        output_path: str,
        df: pd.DataFrame,
        quality_score: float,
        recommendations: List[Dict[str, Any]],
        anomalies: List[Dict[str, Any]]
    ) -> str:
        wb = openpyxl.Workbook()
        
        # Sheet 1: Cleaned Structured Data
        ws_data = wb.active
        ws_data.title = "Structured Data"

        # Headers
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        columns = list(df.columns)
        for col_num, col_name in enumerate(columns, 1):
            cell = ws_data.cell(row=1, column=col_num, value=str(col_name))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        # Data Rows
        for row_num, row_data in enumerate(df.values, 2):
            for col_num, val in enumerate(row_data, 1):
                clean_val = "" if pd.isna(val) else val
                ws_data.cell(row=row_num, column=col_num, value=clean_val)

        # Auto-adjust column widths
        for col in ws_data.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_data.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Sheet 2: Data Quality & Rules Audit
        ws_audit = wb.create_sheet(title="Data Quality Audit")
        ws_audit.cell(row=1, column=1, value="Data Quality Score").font = Font(size=14, bold=True)
        ws_audit.cell(row=1, column=2, value=f"{quality_score:.1f}%").font = Font(size=14, bold=True, color="10B981")

        ws_audit.cell(row=3, column=1, value="Applied Cleaning Rules & Recommendations:").font = Font(size=12, bold=True)
        audit_headers = ["Rule Type", "Target Column", "Description", "Status", "Score Gain"]
        for c_idx, h in enumerate(audit_headers, 1):
            cell = ws_audit.cell(row=4, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill

        for r_idx, rec in enumerate(recommendations, 5):
            ws_audit.cell(row=r_idx, column=1, value=rec.get("rule_type", ""))
            ws_audit.cell(row=r_idx, column=2, value=rec.get("target_column") or "Dataset Wide")
            ws_audit.cell(row=r_idx, column=3, value=rec.get("description", ""))
            ws_audit.cell(row=r_idx, column=4, value=rec.get("status", "accepted"))
            ws_audit.cell(row=r_idx, column=5, value=f"+{rec.get('impact_score_gain', 0):.1f}%")

        # Auto width for audit sheet
        for col in ws_audit.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_audit.column_dimensions[col_letter].width = max(max_len + 3, 15)

        wb.save(output_path)
        return output_path
